"""Rewrites or removes whole paragraphs outside the rebuilt ranges and fixes the abbreviations table.
Usage: python patch_text.py <docx> <patch.json>
patch.json: {"paragraphs": [[prefix, new text or null], ...], "abbr_remove": [...], "abbr_set": {...}, "abbr_add": [[k, v], ...]}"""
import copy, json, sys
import docx

path, patch_path = sys.argv[1:3]
patch = json.load(open(patch_path, encoding="utf-8"))
d = docx.Document(path)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
P = d.paragraphs
start = next(i for i, p in enumerate(P) if p.style.name == "Title1" and p.text.strip() == "LỜI CẢM ƠN")

def set_text(p, text):
    runs = p.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    runs[0].text = text

for prefix, new in patch.get("paragraphs", []):
    hits = [p for p in P[start:] if p.text.strip().startswith(prefix)]
    if len(hits) != 1:
        raise SystemExit(f"{len(hits)} paragraphs start with: {prefix[:60]}")
    if new is None:
        hits[0]._p.getparent().remove(hits[0]._p)
    else:
        set_text(hits[0], new)

for prefix, old, new in patch.get("replace", []):
    hits = [p for p in P[start:] if p.text.strip().startswith(prefix)]
    if len(hits) != 1 or old not in hits[0].text:
        raise SystemExit(f"replace failed: {prefix[:60]}")
    set_text(hits[0], hits[0].text.replace(old, new))

table = d.tables[0]
for row in list(table.rows):
    key = row.cells[0].text.strip()
    if key in patch.get("abbr_remove", []):
        row._tr.getparent().remove(row._tr)
        patch["abbr_remove"].remove(key)
    elif key in patch.get("abbr_set", {}):
        set_text(row.cells[1].paragraphs[0], patch["abbr_set"][key])
template = table.rows[-1]._tr
for key, value in patch.get("abbr_add", []):
    tr = copy.deepcopy(template)
    table._tbl.append(tr)
    row = table.rows[-1]
    set_text(row.cells[0].paragraphs[0], key)
    set_text(row.cells[1].paragraphs[0], value)
d.save(path)
print("patched")
