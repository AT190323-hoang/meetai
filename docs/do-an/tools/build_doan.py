"""Replaces a chapter of the thesis .docx with content written in Markdown, reusing the document's
own paragraphs and tables as formatting templates.
Usage: python build_doan.py <in.docx> <out.docx> <chapter.md> <start heading> <end heading> <image dir>
"""
import copy
import re
import sys

import docx
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docx.table import Table

import json
if sys.argv[1] == "--spec":
    spec = json.load(open(sys.argv[2], encoding="utf-8"))
    src, out, img_dir, RANGES = spec["src"], spec["out"], spec["img_dir"], spec["ranges"]
else:
    src, out, md_path, start_text, end_text, img_dir = sys.argv[1:7]
    RANGES = [{"md": md_path, "start": start_text, "end": end_text}]
d = docx.Document(src)
P = d.paragraphs
body = d.element.body

def body_paragraphs():
    ps = d.paragraphs
    first = next(i for i, p in enumerate(ps) if p.style.name == "Title1" and p.text.strip() == "LỜI CẢM ƠN")
    return ps[first:]

def find(text):
    for p in body_paragraphs():
        if p.text.strip() == text:
            return p
    for p in body_paragraphs():
        if p.text.strip().startswith(text):
            return p
    raise SystemExit(f"heading not found: {text}")

def keep_elements(tag):
    """The caption paragraph named `tag` (e.g. "Bảng 1.2") and the table or image it belongs to."""
    cap = next(p for p in body_paragraphs() if p.style.name in ("bảng", "hình") and p.text.strip().startswith(tag + ":"))
    prev = cap._p.getprevious()
    while prev is not None and prev.tag != W + "tbl" and not prev.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
        prev = prev.getprevious()
    return [prev, cap._p]

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

def tpl(style, pred=lambda p: True):
    for p in P:
        if p.style.name == style and p.text.strip() and pred(p):
            return copy.deepcopy(p._p)
    raise SystemExit(f"template not found: {style}")

# Templates taken from the current Chapter 3 and Chapter 2
T = {
    "h2": tpl("Title2", lambda p: re.match(r"\d\.\d\. ", p.text)),
    "h3": tpl("Title3", lambda p: re.match(r"\d\.\d\.\d\. ", p.text)),
    "body": tpl("Normal (Web)", lambda p: p.paragraph_format.alignment == 3 and len(p.text) > 200 and not p._p.xpath(".//w:numPr")),
    "bullet": tpl("List Paragraph", lambda p: p.text.startswith("Cấu hình A")),
    "cap_table": tpl("bảng"),
    "cap_fig": tpl("hình"),
    "code": tpl("Normal (Web)", lambda p: any(r.font.name == "Courier New" for r in p.runs)),
}
fig_para = next(p for p in P if p._p.xpath(".//w:drawing") and p.paragraph_format.alignment == 1)
T["figure"] = copy.deepcopy(fig_para._p)
table_tpl = next(t for t in d.tables if t.cell(0, 0).text.startswith("Thành phần"))

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

def strip_runs(p_el):
    for child in list(p_el):
        if child.tag != W + "pPr":
            p_el.remove(child)

def first_rpr(template):
    r = template.find(W + "r")
    rpr = r.find(W + "rPr") if r is not None else None
    return copy.deepcopy(rpr) if rpr is not None else None

def add_runs(para, text, base_rpr, code_font=False):
    # **bold** and `code` inline markup
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        bold = part.startswith("**")
        code = part.startswith("`")
        content = part[2:-2] if bold else part[1:-1] if code else part
        run = para.add_run(content)
        if base_rpr is not None:
            run._r.insert(0, copy.deepcopy(base_rpr))
        if bold:
            run.bold = True
        if code or code_font:
            run.font.name = "Courier New"

def make_para(kind, text, parent):
    el = copy.deepcopy(T[kind])
    rpr = first_rpr(T[kind])
    strip_runs(el)
    para = Paragraph(el, parent)
    add_runs(para, text, rpr, code_font=(kind == "code"))
    if kind == "code":
        fmt = para.paragraph_format
        fmt.first_line_indent = Cm(0)
        fmt.left_indent = Cm(0.75)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        fmt.alignment = 0
        for run in para.runs:
            run.font.size = Pt(10)
    return el

def make_figure(path, parent, width_cm):
    el = copy.deepcopy(T["figure"])
    strip_runs(el)
    para = Paragraph(el, parent)
    para.add_run().add_picture(path, width=Cm(width_cm))
    return el

def make_table(rows, parent, percents=None):
    el = copy.deepcopy(table_tpl._tbl)
    tbl = Table(el, parent)
    header_tpl = copy.deepcopy(tbl.rows[0]._tr)
    body_tpl = copy.deepcopy(tbl.rows[1]._tr)
    for tr in list(el.findall(W + "tr")):
        el.remove(tr)
    ncols = len(rows[0])
    # Column widths proportional to content length, total 9060 dxa
    # Each column fits its longest word (about 115 dxa per character at 11 pt), the rest of the
    # 9060 dxa is shared by content length
    plain = lambda t: re.sub(r"[*`]", "", t)
    mins = [max(max(len(w) for r in rows for w in (plain(r[i]).split() or [""])) * 115 + 250, 700) for i in range(ncols)]
    lens = [max(min(len(plain(r[i])), 90) for r in rows) + 4 for i in range(ncols)]
    spare = max(9060 - sum(mins), 0)
    widths = [int(mins[i] + spare * lens[i] / sum(lens)) for i in range(ncols)]
    scale = 9060 / sum(widths)
    widths = [int(w * scale) for w in widths]
    widths[-1] += 9060 - sum(widths)
    if percents:
        widths = [int(9060 * x / sum(percents)) for x in percents]
        widths[-1] += 9060 - sum(widths)
    grid = el.find(W + "tblGrid")
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        gc = copy.deepcopy(T_GRIDCOL)
        gc.set(W + "w", str(w))
        grid.append(gc)
    for ri, row in enumerate(rows):
        tr = copy.deepcopy(header_tpl if ri == 0 else body_tpl)
        tcs = tr.findall(W + "tc")
        while len(tcs) < ncols:
            tr.append(copy.deepcopy(tcs[-1]))
            tcs = tr.findall(W + "tc")
        for extra in tcs[ncols:]:
            tr.remove(extra)
        for ci, tc in enumerate(tr.findall(W + "tc")):
            tc.find(W + "tcPr").find(W + "tcW").set(W + "w", str(widths[ci]))
            ps = tc.findall(W + "p")
            for extra in ps[1:]:
                tc.remove(extra)
            p_el = ps[0]
            rpr = first_rpr(p_el)
            strip_runs(p_el)
            add_runs(Paragraph(p_el, parent), row[ci], rpr)
            if ri == 0:
                for r in Paragraph(p_el, parent).runs:
                    r.bold = True
        el.append(tr)
    return el

T_GRIDCOL = copy.deepcopy(table_tpl._tbl.find(W + "tblGrid")[0])

# ---------------------------------------------------------------- parse Markdown into blocks, replace a range
def apply_range(md_path, start_text, end_text):
    lines = open(md_path, encoding="utf-8").read().split("\n")
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            j = i + 1
            while not lines[j].startswith("```"):
                j += 1
            blocks.append(("code", lines[i + 1:j]))
            i = j + 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        s = line.strip()
        m_k = re.match(r"\{\{keep: (.+?)\}\}$", s)
        m_w = re.match(r"<!-- widths: ([\d, ]+) -->", s)
        if m_k:
            blocks.append(("keep", m_k.group(1).strip()))
        elif m_w:
            blocks.append(("widths", [int(x) for x in m_w.group(1).split(",")]))
        elif not s or s.startswith(">") or s.startswith("# "):
            pass
        elif s.startswith("### "):
            blocks.append(("h3", s[4:]))
        elif s.startswith("## "):
            blocks.append(("h2", s[3:]))
        elif s.startswith("- "):
            blocks.append(("bullet", s[2:]))
        elif re.match(r"Bảng \d+\.\d+: ", s):
            blocks.append(("cap_table", s))
        elif re.match(r"Hình \d+\.\d+: .*\[.+\]$", s):
            m = re.match(r"(Hình [\d.]+: .*?) \[(.+?)\]$", s)
            blocks.append(("figure", (m.group(1), m.group(2))))
        else:
            blocks.append(("body", s))
        i += 1

    # Captions in the thesis sit below their table: move each "Bảng" caption after the table it names
    ordered = []
    k = 0
    while k < len(blocks):
        if (blocks[k][0] == "cap_table" and k + 2 < len(blocks) and blocks[k + 1][0] == "widths"
                and blocks[k + 2][0] == "table"):
            ordered += [blocks[k + 1], blocks[k + 2], blocks[k]]
            k += 3
        elif blocks[k][0] == "cap_table" and k + 1 < len(blocks) and blocks[k + 1][0] == "table":
            ordered += [blocks[k + 1], blocks[k]]
            k += 2
        else:
            ordered.append(blocks[k])
            k += 1

    # ---------------------------------------------------------------- replace the chapter
    kept = {content: keep_elements(content) for kind, content in ordered if kind == "keep"}
    start = find(start_text)._p
    end = find(end_text)._p
    children = list(body)
    s_idx, e_idx = children.index(start), children.index(end)
    kept_tail = []
    for el in children[s_idx + 1:e_idx]:
        # Keep section breaks and page breaks that end the chapter
        has_sect = el.find(".//" + W + "sectPr") is not None
        page_break = any(br.get(W + "type") == "page" for br in el.iter(W + "br"))
        if has_sect or (el.tag == W + "p" and page_break and not "".join(el.itertext()).strip()):
            kept_tail.append(el)
            continue
        body.remove(el)

    anchor = kept_tail[0] if kept_tail else end
    parent = find(end_text)._parent
    next_widths = None
    for kind, content in ordered:
        if kind == "keep":
            for el in kept[content]:
                anchor.addprevious(el)
        elif kind == "widths":
            next_widths = content
        elif kind == "table":
            anchor.addprevious(make_table(content, parent, next_widths))
            next_widths = None
        elif kind == "code":
            for code_line in content:
                anchor.addprevious(make_para("code", code_line if code_line else " ", parent))
        elif kind == "figure":
            caption, file = content
            path = f"docs/do-an/{file[1:]}" if file.startswith("@") else f"{img_dir}/{file}"
            width = 13.0 if file.endswith(".png") else 15.5
            anchor.addprevious(make_figure(path, parent, width))
            anchor.addprevious(make_para("cap_fig", caption, parent))
        else:
            anchor.addprevious(make_para(kind, content, parent))


for r in RANGES:
    apply_range(r["md"], r["start"], r["end"])

d.save(out)
print("saved", out)
