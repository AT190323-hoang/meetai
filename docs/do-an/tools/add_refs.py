import copy, sys, docx
from docx.text.paragraph import Paragraph
path = sys.argv[1]
d = docx.Document(path)
last = [p for p in d.paragraphs if p.text.startswith("[25] ")][-1]
refs = [
    "[26] Chao, P., Debenedetti, E., Robey, A., Andriushchenko, M., Croce, F., Sehwag, V., Dobriban, E., Flammarion, N., Pappas, G. J., Tramèr, F., Hassani, H., and Wong, E. JailbreakBench: An Open Robustness Benchmark for Jailbreaking Large Language Models. NeurIPS Datasets and Benchmarks Track, 2024.",
    "[27] Chao, P., Robey, A., Dobriban, E., Hassani, H., Pappas, G. J., and Wong, E. Jailbreaking Black Box Large Language Models in Twenty Queries. arXiv:2310.08419, 2023.",
    "[28] Quốc hội nước Cộng hòa xã hội chủ nghĩa Việt Nam. Luật An ninh mạng, Luật số 24/2018/QH14, ngày 12/6/2018.",
]
anchor = last._p
for text in refs:
    el = copy.deepcopy(last._p)
    rpr = None
    for r in el.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r"):
        rp = r.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
        if rp is not None and rp.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}i") is None:
            rpr = copy.deepcopy(rp); break
    for child in list(el):
        if not child.tag.endswith("}pPr"):
            el.remove(child)
    para = Paragraph(el, last._parent)
    run = para.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    anchor.addnext(el)
    anchor = el
d.save(path)
print("refs added")
